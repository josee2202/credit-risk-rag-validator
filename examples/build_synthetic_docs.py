"""Genera los documentos sintéticos de examples/synthetic/ a partir de synthetic_corpus.py.

    python examples/build_synthetic_docs.py

Requiere fpdf2 (requirements-dev.txt). Los ficheros generados ya están versionados; solo hace
falta ejecutarlo si se modifica el corpus.
"""
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fpdf import FPDF

import synthetic_corpus as c

OUT = Path(__file__).resolve().parent / "synthetic"


def build_normativa_pdf(path: Path):
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 20, 20)
    pdf.set_creation_date(datetime(2026, 1, 1, tzinfo=timezone.utc))  # fecha fija: PDF reproducible
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.multi_cell(0, 8, c.TITULO_NORMATIVA, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font("Helvetica", "I", 10)
    pdf.multi_cell(0, 5, c.INTRO_NORMATIVA, new_x="LMARGIN", new_y="NEXT")
    for rid, (titulo, texto) in c.REQUISITOS.items():
        pdf.ln(5)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 6, f"{rid}. {titulo}", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, texto, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def build_docx(path: Path, titulo: str, intro: str, secciones):
    doc = Document()
    doc.core_properties.author = "Documento sintético"
    doc.add_heading(titulo, level=1)
    doc.add_paragraph(intro).italic = True
    for i, (sub, texto) in enumerate(secciones, 1):
        doc.add_heading(f"{i}. {sub}", level=2)
        doc.add_paragraph(texto)
    doc.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    build_normativa_pdf(OUT / "normativa_sintetica_pd.pdf")
    build_docx(OUT / "metodologia_banco_ficticio.docx", c.TITULO_METODOLOGIA, c.INTRO_METODOLOGIA,
               [(s["titulo"], s["texto"]) for s in c.METODOLOGIA])
    build_docx(OUT / "tecnico_banco_ficticio.docx", c.TITULO_TECNICO, c.INTRO_TECNICO, c.TECNICO)
    (OUT / "codigo_pd_ficticio.sas").write_text(c.CODIGO_SAS, encoding="utf-8")
    print("Documentos generados en", OUT)
